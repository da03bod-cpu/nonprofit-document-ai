import json
from pathlib import Path
from docx import Document


def parse_docx(file_path: str):
    doc = Document(file_path)

    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    # Include table content because annual reports often store
    # important program/KPI data inside tables.
    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"===== TABLE {table_index} =====")

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def parse_json(file_path: str):
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Preserve the original JSON semantics, but give Qwen a readable
    # representation.
    return json.dumps(data, ensure_ascii=False, indent=2)


def parse_document(file_path: str, extension: str):
    extension = extension.lower().lstrip(".")

    if extension == "docx":
        return parse_docx(file_path)

    if extension == "json":
        return parse_json(file_path)

    raise ValueError(f"Unsupported direct document type: {extension}")
